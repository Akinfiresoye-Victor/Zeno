"""
Document generation tool for docx, md, and txt documents.
Category: file_write
"""

from __future__ import annotations

import re
from pathlib import Path

from zeno.tools.file_tools import validate_and_resolve_path

try:
    import docx
except ImportError:
    docx = None


def sanitize_filename(name: str) -> str:
    """Strip unsafe filesystem characters from title for clean filename."""
    cleaned = re.sub(r'[\\/*?:"<>|]', "", name).strip()
    cleaned = re.sub(r"\s+", "_", cleaned)
    if not cleaned:
        cleaned = "untitled_document"
    return cleaned[:100]


def create_document(
    title: str,
    content: str,
    file_type: str = "docx",
    folder: str = "Documents",
) -> str:
    """
    Create a formatted document (.docx, .md, or .txt) in an allowed folder.

    Args:
        title: Document heading and base filename.
        content: Document body text. Paragraphs are split on newlines.
        file_type: Document format, one of 'docx', 'md', or 'txt'. Defaults to 'docx'.
        folder: Subdirectory inside allowed folders (e.g. 'Documents').

    Returns:
        str: Absolute filesystem path to the created document.
    """
    fmt = file_type.lower().strip().lstrip(".")
    if fmt not in ["docx", "md", "txt"]:
        raise ValueError("file_type must be one of: 'docx', 'md', or 'txt'.")

    # Target folder inside allowed folders
    target_dir = validate_and_resolve_path(folder)
    target_dir.mkdir(parents=True, exist_ok=True)

    filename = f"{sanitize_filename(title)}.{fmt}"
    file_path = (target_dir / filename).resolve()

    if fmt == "docx":
        if docx is None:
            raise RuntimeError("python-docx is not installed. Install with: pip install python-docx==1.1.2")
        doc = docx.Document()
        doc.add_heading(title, level=0)
        paragraphs = content.split("\n")
        for para in paragraphs:
            para_str = para.strip()
            if para_str:
                doc.add_paragraph(para_str)
        doc.save(str(file_path))

    elif fmt == "md":
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(f"# {title}\n\n{content}\n")

    else:  # txt
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(f"{title}\n{'=' * len(title)}\n\n{content}\n")

    return str(file_path)
